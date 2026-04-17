from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def convert_links(text: str) -> str:
    return LINK_RE.sub(r"\1 (\2)", text)


def set_default_style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)


def add_body_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.first_line_indent = Cm(0.7)
    para.add_run(convert_links(text.strip()))


def add_reference_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.left_indent = Cm(1.0)
    para.paragraph_format.first_line_indent = Cm(-1.0)
    para.add_run(convert_links(text.strip()))


def add_figure(doc: Document, image_path: Path, caption: str, width_cm: float = 16.0) -> None:
    if not image_path.exists():
        raise FileNotFoundError(f"Image not found: {image_path}")

    image_para = doc.add_paragraph()
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_para.paragraph_format.space_after = Pt(4)
    image_run = image_para.add_run()
    image_run.add_picture(str(image_path), width=Cm(width_cm))

    if caption.strip():
        caption_para = doc.add_paragraph()
        caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_para.paragraph_format.space_after = Pt(8)
        caption_run = caption_para.add_run(convert_links(caption.strip()))
        caption_run.italic = True
        caption_run.font.name = "Times New Roman"
        caption_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        caption_run.font.size = Pt(10)


def add_heading(doc: Document, text: str, level: int) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run(text.strip())
    run.bold = True
    run.font.name = "Times New Roman"
    if level == 1:
        run.font.size = Pt(16)
    else:
        run.font.size = Pt(13)


def add_title(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(12)
    run = para.add_run(text.strip())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)


def split_table_row(line: str) -> List[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def is_table_separator(line: str) -> bool:
    stripped = line.replace("|", "").replace("-", "").replace(":", "").strip()
    return stripped == ""


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("left", "top", "right", "bottom"):
        edge_data = kwargs.get(edge)
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        if edge_data:
            for key, value in edge_data.items():
                element.set(qn(f"w:{key}"), str(value))
        else:
            element.set(qn("w:val"), "nil")


def apply_three_line_table(table, body_row_count: int, col_count: int) -> None:
    border_on = {"val": "single", "sz": 8, "color": "000000", "space": 0}
    border_off = {"val": "nil"}

    total_rows = 1 + body_row_count
    for r in range(total_rows):
        for c in range(col_count):
            cell = table.cell(r, c)
            top = border_off
            bottom = border_off
            if r == 0:
                top = border_on
                bottom = border_on
            elif r == total_rows - 1:
                bottom = border_on

            set_cell_border(
                cell,
                top=top,
                bottom=bottom,
                left=border_off,
                right=border_off,
            )


def add_table(doc: Document, table_lines: List[str]) -> None:
    rows = [split_table_row(line) for line in table_lines if line.strip()]
    if len(rows) < 2:
        return

    header = rows[0]
    body = [row for row in rows[1:] if not is_table_separator("|" + "|".join(row) + "|")]

    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, cell_text in enumerate(header):
        table.cell(0, i).text = convert_links(cell_text)
        for paragraph in table.cell(0, i).paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10.5)

    for r, row in enumerate(body, start=1):
        padded = row + [""] * (len(header) - len(row))
        for c, cell_text in enumerate(padded[: len(header)]):
            table.cell(r, c).text = convert_links(cell_text)
            for paragraph in table.cell(r, c).paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)

    apply_three_line_table(table, len(body), len(header))


def markdown_to_docx(input_path: Path, output_path: Path) -> None:
    lines = input_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    set_default_style(doc)

    paragraph_buffer: List[str] = []
    table_buffer: List[str] = []
    in_references = False

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_body_paragraph(doc, " ".join(part.strip() for part in paragraph_buffer if part.strip()))
            paragraph_buffer = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_table(doc, table_buffer)
            table_buffer = []

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("|"):
            flush_paragraph()
            table_buffer.append(line)
            continue

        flush_table()

        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            add_title(doc, stripped[2:])
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            heading_text = stripped[3:]
            add_heading(doc, heading_text, level=1)
            in_references = heading_text.lower() == "references"
            continue

        image_match = IMAGE_RE.fullmatch(stripped)
        if image_match:
            flush_paragraph()
            caption, image_ref = image_match.groups()
            image_path = Path(image_ref)
            if not image_path.is_absolute():
                image_path = (input_path.parent / image_path).resolve()
            add_figure(doc, image_path, caption)
            continue

        if in_references:
            flush_paragraph()
            add_reference_paragraph(doc, stripped)
            continue

        if re.match(r"^\d+\.\s", stripped):
            flush_paragraph()
            add_body_paragraph(doc, stripped)
        else:
            paragraph_buffer.append(stripped)

    flush_table()
    flush_paragraph()
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python export_markdown_to_docx.py <input.md> <output.docx>")

    input_file = Path(sys.argv[1]).expanduser().resolve()
    output_file = Path(sys.argv[2]).expanduser().resolve()
    markdown_to_docx(input_file, output_file)
    print(f"Exported: {output_file}")
