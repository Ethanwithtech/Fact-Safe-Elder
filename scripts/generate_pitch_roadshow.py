"""
生成 FactSafe 竞赛 6 分钟路演 Word 稿（含真实案例复盘与技术亮点）
输出: FactSafe_竞赛路演方案_6分钟.docx
"""
import os
from datetime import date

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "FactSafe_竞赛路演方案_6分钟.docx")
MD_SRC = os.path.join(ROOT, "FactSafe_6分钟Demo演讲稿.md")


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    for r in p.runs:
        r.font.size = Pt(11)


def main():
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("FactSafe · 6 分钟竞赛 Demo 演讲稿\n（真实案例 + 技术突破 + 诚实边界）")
    tr.bold = True
    tr.font.size = Pt(18)
    doc.add_paragraph()
    add_para(doc, f"更新日期：{date.today().isoformat()}  |  系统：localhost:8000/health 需 all_models_ready=true")

    add_heading(doc, "一、答辩核心逻辑", 1)
    add_para(doc, "用「平台结构性盲区 → 我们消费端补位 → 实测数据 → 社会价值」闭环。不说「比字节强」，说「平台商业逻辑做不了我们这一层」。")

    add_heading(doc, "二、技术亮点与四大突破", 1)
    breakthroughs = [
        ("突破① 跨模态联合推理", "字幕合规+口播诈骗", "OCR/ASR 分通道+散度+mismatch", "18条错位集联合检出率100%；王阿姨案 joint_risk=1.0"),
        ("突破② 认知多任务模型", "老人专属话术", "1799条标注+5类手法 MacBERT", "risk F1≈0.93；authority_forgery 等 AI 标签"),
        ("突破③ 流式增量推理", "发布端只审一次", "SSE progress_risk+风险轨迹", "观看过程中分数持续上升"),
        ("突破④ 证据链可解释", "黑盒不可信", "evidence_chain+案例RAG+GPT纠正", "右侧面板+飞书卡片推送"),
    ]
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(["突破", "平台痛点", "方案", "数据/演示"]):
        t.rows[0].cells[i].text = h
    for row in breakthroughs:
        c = t.add_row().cells
        for i, v in enumerate(row):
            c[i].text = v

    add_heading(doc, "三、6 分钟逐秒脚本", 1)
    script = [
        ("0:00–0:40", "开场", "四件平台做不到的事：跨模态阴阳、老人话术、观看实时窗、黑盒解释。FactSafe=消费端补充层。"),
        ("0:40–1:30", "王阿姨案", "浦东72岁、字幕养生+语音秘方加微信、12万赞零拦截、48万损失。点出「阴阳视频」。"),
        ("1:30–3:30", "演示", "localhost:3000 → 灵动岛多留心 → 跨模态条 → 手法AI标签 → 风险轨迹 → 飞书通知 → 跳过不看。"),
        ("3:30–4:40", "数据", "错位集100%；认知模型已加载；流式SSE；证据链面板。"),
        ("4:40–5:20", "诚实", "私域微信/银行转账不能替代；数据1799条；原型=上传/模拟器，量产走无障碍OCR。"),
        ("5:20–6:00", "收尾", "1.6亿老人、39万起案。一句话：消费端可解释可通知的老年短视频反诈原型。"),
    ]
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = "Table Grid"
    t2.rows[0].cells[0].text = "时间"
    t2.rows[0].cells[1].text = "环节"
    t2.rows[0].cells[2].text = "内容"
    for a, b, c in script:
        r = t2.add_row().cells
        r[0].text, r[1].text, r[2].text = a, b, c

    add_heading(doc, "四、王阿姨案 — 系统能否解决？", 1)
    add_para(doc, "结论：短视频引流第1天可拦截（已实测）；微信私域与转账需家人协同，不能夸大。", bold=True)
    mapping = [
        ("2/10 短视频", "平台单模态漏检", "能", "跨模态 mismatch + 飞书告警 + 适老化提醒"),
        ("2/11 加微信", "私域不可审", "间接", "口播「加微信」高风险；提醒子女关注"),
        ("2/12–25 情感洗脑", "微信一对一", "不能自动", "文本可检测情感话术，非实时监听微信"),
        ("2/25 转账", "银行泛化短信", "间接", "促子女介入，不冻结账户"),
    ]
    t3 = doc.add_table(rows=1, cols=4)
    t3.style = "Table Grid"
    for i, h in enumerate(["阶段", "失效原因", "FactSafe", "说明"]):
        t3.rows[0].cells[i].text = h
    for row in mapping:
        c = t3.add_row().cells
        for i, v in enumerate(row):
            c[i].text = v

    add_heading(doc, "五、Demo 测试文案（王阿姨阴阳视频）", 1)
    add_para(doc, "OCR：糖尿病日常调理三招，管住嘴迈开腿，保持好心情")
    add_para(doc, "ASR：祖传秘方…加我屏幕下方的微信…")
    add_para(doc, "实测(2026-06-03)：mismatch=true, joint_risk=1.0, danger≈0.85, 手法 authority_forgery, 来源 cognitive_model")

    add_heading(doc, "六、竞赛演示模式（必开）", 1)
    add_bullet(doc, "顶栏开启「竞赛演示模式」→ 70% 手机 + 30% 检测一屏展示")
    add_bullet(doc, "自动：隐藏设置/统计、锁定宁漏勿误、预加载模型与 3 条 Demo 视频、同步飞书")
    add_bullet(doc, "点击「保健品诈骗 / 金融诈骗 / 安全科普」一键播片并触发检测")

    add_heading(doc, "七、上场检查清单", 1)
    for c in [
        "竞赛演示模式 ON，横幅显示「演示资源就绪」",
        "GET /health → all_models_ready: true（三模型均已加载）",
        "飞书群能收到危险视频告警卡片",
        "完整 Markdown 稿：FactSafe_6分钟Demo演讲稿.md",
    ]:
        add_bullet(doc, c)

    doc.save(OUT)
    print(f"已生成: {OUT}")
    if os.path.isfile(MD_SRC):
        print(f"Markdown 稿: {MD_SRC}")


if __name__ == "__main__":
    main()
